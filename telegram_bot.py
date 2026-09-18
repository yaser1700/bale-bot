import os
import re
import time
import threading
import requests
from flask import Flask, request, jsonify
from openpyxl import load_workbook
from urllib.parse import unquote

# =========================================================
# CONFIG
# =========================================================

TOKEN = os.getenv("TELEGRAM_TOKEN", "").strip()
ADMIN_CHAT_ID = os.getenv("TELEGRAM_ADMIN_ID", "").strip()

WEBSITE = "https://www.tecnoyadakabbasi.ir"
PHONE = "09377700031"

BASE_URL = f"https://api.telegram.org/bot{TOKEN}" if TOKEN else ""

app = Flask(__name__)

CARTS = {}
USER_STATES = {}
LOCK = threading.RLock()
ORDER_COUNTER = 1000

PDF_GROUPS = {
    "📦 سوکت عباسی": "سوکت عباسی",
    "🔌 کابل تکنو سبزوار": "کابل تکنو سبزوار",
    "⚡ وایر عباسی": "وایر عباسی",
    "🔩 مهره و سنسور": "مهره و سنسور",
    "💡 قطعات برقی خودرو": "قطعات برقی خودرو",
    "🧩 خارجات و پلیمریجات": "خارجات و پلیمریجات",
    "🔌 کابل خودرو سبزوار": "کابل خودرو سبزوار",
    "⚙️ شیلنگ خودرو": "شیلنگ خودرو",
    "⚙️ جلوبندی": "جلوبندی",
}

# =========================================================
# HELPERS
# =========================================================

def normalize(text):
    text = str(text or "")
    text = text.translate(str.maketrans(
        "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
        "01234567890123456789"
    ))
    replacements = {
        "ي": "ی", "ى": "ی", "ك": "ک", "ۀ": "ه", "ة": "ه",
        "ؤ": "و", "إ": "ا", "أ": "ا", "\u200c": " ",
        "\u200f": "", "\u200e": ""
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def compact(text):
    return re.sub(r"[^0-9a-zآ-ی]+", "", normalize(text))


def tg_request(method, data=None, files=None, timeout=30):
    if not TOKEN:
        print("TELEGRAM_TOKEN NOT FOUND")
        return None
    try:
        return requests.post(
            f"{BASE_URL}/{method}",
            data=data if files else None,
            json=data if not files else None,
            files=files,
            timeout=timeout,
        )
    except Exception as e:
        print("TELEGRAM API ERROR:", repr(e))
        return None


def send_message(chat_id, text, keyboard=None, inline=False):
    data = {"chat_id": chat_id, "text": str(text)}
    if keyboard:
        if inline:
            data["reply_markup"] = {"inline_keyboard": keyboard}
        else:
            data["reply_markup"] = {
                "keyboard": keyboard,
                "resize_keyboard": True,
            }
    response = tg_request("sendMessage", data=data)
    if response is not None:
        print("SEND MESSAGE:", response.status_code, response.text[:300])
    return response


def answer_callback(callback_id, text=""):
    if not callback_id:
        return
    response = tg_request(
        "answerCallbackQuery",
        data={"callback_query_id": callback_id, "text": text},
    )
    if response is not None:
        print("CALLBACK:", response.status_code, response.text[:200])


def main_menu(chat_id):
    keyboard = [
        ["📞 تماس مستقیم", "🌐 ورود به سایت"],
        ["📄 دریافت لیست قیمت"],
        ["🔎 جستجوی کالا"],
        ["🛒 سبد خرید", "🧾 ثبت سفارش"],
        ["📦 سوکت عباسی", "🔌 کابل تکنو سبزوار"],
        ["⚡ وایر عباسی", "🔩 مهره و سنسور"],
        ["💡 قطعات برقی خودرو", "🧩 خارجات و پلیمریجات"],
        ["🔌 کابل خودرو سبزوار", "⚙️ شیلنگ خودرو"],
        ["⚙️ جلوبندی"],
    ]
    send_message(chat_id, "📋 گزینه مورد نظر را انتخاب کنید:", keyboard)


def find_pdf(group_name):
    base = os.path.dirname(os.path.abspath(__file__))
    wanted = compact(group_name)
    try:
        filenames = os.listdir(base)
    except Exception as e:
        print("LIST FILE ERROR:", repr(e))
        return None

    for filename in filenames:
        if not filename.lower().endswith(".pdf"):
            continue
        decoded = unquote(filename)
        name_without_ext = os.path.splitext(decoded)[0]
        normalized_name = compact(name_without_ext)
        if wanted in normalized_name or normalized_name in wanted:
            return os.path.join(base, filename)
    return None


def send_pdf(chat_id, pdf_path, caption):
    try:
        with open(pdf_path, "rb") as file:
            response = requests.post(
                f"{BASE_URL}/sendDocument",
                data={"chat_id": str(chat_id), "caption": caption},
                files={
                    "document": (
                        os.path.basename(pdf_path),
                        file,
                        "application/pdf",
                    )
                },
                timeout=180,
            )
        print("SEND PDF:", response.status_code, response.text[:300])
        return response
    except Exception as e:
        print("PDF ERROR:", repr(e))
        return None


def format_price(value):
    if value is None:
        return ""
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, float):
        if value.is_integer():
            return f"{int(value):,}"
        return str(value)
    text = str(value).strip().replace(",", "").replace("٬", "")
    try:
        number = float(text)
        if number.is_integer():
            return f"{int(number):,}"
    except Exception:
        pass
    return text


def price_number(value):
    try:
        text = normalize(value).replace(",", "").replace("٬", "")
        return int(float(text))
    except Exception:
        return 0


def get_columns(row):
    columns = {}
    for index, value in enumerate(row):
        if value is None:
            continue
        name = normalize(value)
        if "گروه" in name:
            columns["group"] = index
        elif "کد کالا" in name or name == "کد" or "شناسه" in name:
            columns["code"] = index
        elif "نام کالا" in name or name == "نام":
            columns["name"] = index
        elif "قیمت" in name:
            columns["price"] = index
        elif "توضیحات" in name or "توضیح" in name:
            columns["description"] = index
    return columns


def search_matches(query, code, name, group, description):
    q = normalize(query)
    qc = compact(query)
    if not q:
        return False

    full_text = " ".join([
        normalize(code),
        normalize(name),
        normalize(group),
        normalize(description),
    ])
    full_compact = compact(full_text)

    if q in full_text or (qc and qc in full_compact):
        return True

    words = [word for word in q.split() if len(word) >= 2]
    return bool(words) and all(word in full_text for word in words)


def search_excel(query):
    base = os.path.dirname(os.path.abspath(__file__))
    try:
        files = os.listdir(base)
    except Exception as e:
        print("FILES ERROR:", repr(e))
        return []

    excel_files = [
        f for f in files
        if f.lower().endswith(".xlsx") and not f.startswith("~$")
    ]

    results = []
    seen = set()

    for filename in excel_files:
        path = os.path.join(base, filename)
        workbook = None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)

            for sheet in workbook.worksheets:
                header_row = None
                columns = None

                for row_number, row in enumerate(
                    sheet.iter_rows(values_only=True), start=1
                ):
                    found = get_columns(row)
                    if "code" in found and "name" in found and "price" in found:
                        header_row = row_number
                        columns = found
                        break

                if not columns:
                    continue

                for row in sheet.iter_rows(
                    min_row=header_row + 1, values_only=True
                ):
                    if not row:
                        continue

                    def get_value(key):
                        index = columns.get(key)
                        if index is None or index >= len(row):
                            return ""
                        return str(row[index] or "").strip()

                    group = get_value("group")
                    code = get_value("code")
                    name = get_value("name")
                    description = get_value("description")

                    price_index = columns.get("price")
                    price = ""
                    if price_index is not None and price_index < len(row):
                        price = format_price(row[price_index])

                    if not code and not name:
                        continue

                    if not search_matches(query, code, name, group, description):
                        continue

                    key = (
                        normalize(code),
                        normalize(name),
                        normalize(group),
                        price,
                    )
                    if key in seen:
                        continue
                    seen.add(key)

                    results.append({
                        "code": code,
                        "name": name,
                        "price": price,
                        "group": group,
                        "description": description,
                    })

            workbook.close()

        except Exception as e:
            print("EXCEL ERROR:", filename, repr(e))
            try:
                if workbook:
                    workbook.close()
            except Exception:
                pass

    print("SEARCH:", query, "RESULT COUNT:", len(results))
    return results


def send_search_results(chat_id, results):
    if not results:
        send_message(chat_id, "❌ کالایی با این کد یا نام پیدا نشد.")
        return

    USER_STATES[chat_id] = {
        "state": "search_results",
        "results": results,
    }

    send_message(
        chat_id,
        f"🔎 {len(results)} کالا پیدا شد:\n\n"
        "برای هر کالا می‌توانید «افزودن به سبد خرید» را بزنید.",
    )

    for index, item in enumerate(results):
        message = (
            f"📦 کد کالا: {item['code']}\n"
            f"📝 نام کالا: {item['name']}\n"
            f"💰 قیمت: {item['price']} ریال\n"
            f"📁 گروه: {item['group']}"
        )
        if item["description"]:
            message += f"\nℹ️ توضیحات: {item['description']}"

        keyboard = [[{
            "text": "➕ افزودن به سبد خرید",
            "callback_data": f"ADD:{index}",
        }]]

        send_message(chat_id, message, keyboard, inline=True)


def add_to_cart(chat_id, item, quantity):
    with LOCK:
        cart = CARTS.setdefault(chat_id, [])
        item_code = normalize(item["code"])

        for cart_item in cart:
            if normalize(cart_item["code"]) == item_code:
                cart_item["quantity"] += quantity
                return

        cart.append({
            "code": item["code"],
            "name": item["name"],
            "price": price_number(item["price"]),
            "quantity": quantity,
        })


def show_cart(chat_id):
    with LOCK:
        cart = list(CARTS.get(chat_id, []))

    if not cart:
        send_message(chat_id, "🛒 سبد خرید شما خالی است.")
        return

    message = "🛒 سبد خرید شما:\n\n"
    total = 0

    for index, item in enumerate(cart, start=1):
        item_total = item["price"] * item["quantity"]
        total += item_total
        message += (
            f"{index}. {item['name']}\n"
            f"🔢 کد: {item['code']}\n"
            f"📦 تعداد: {item['quantity']}\n"
            f"💰 قیمت واحد: {item['price']:,} ریال\n"
            f"💵 مبلغ: {item_total:,} ریال\n"
            "──────────────\n"
        )

    message += f"\n💰 جمع کل: {total:,} ریال"

    keyboard = [
        ["🧾 ثبت سفارش"],
        ["🗑️ خالی کردن سبد"],
        ["🔎 جستجوی کالا"],
        ["🔙 منوی اصلی"],
    ]
    send_message(chat_id, message, keyboard)


def clear_cart(chat_id):
    with LOCK:
        CARTS[chat_id] = []
        USER_STATES.pop(chat_id, None)
    send_message(chat_id, "🗑️ سبد خرید خالی شد.")
    main_menu(chat_id)


def start_order(chat_id):
    with LOCK:
        cart = list(CARTS.get(chat_id, []))

    if not cart:
        send_message(
            chat_id,
            "🛒 سبد خرید خالی است.\nابتدا کالا به سبد اضافه کنید.",
        )
        return

    USER_STATES[chat_id] = {"state": "order_name"}
    send_message(
        chat_id,
        "🧾 ثبت سفارش\n\nلطفاً نام و نام خانوادگی خود را وارد کنید:",
    )


def finish_order(chat_id, name, phone):
    global ORDER_COUNTER

    with LOCK:
        cart = list(CARTS.get(chat_id, []))
        if not cart:
            send_message(chat_id, "❌ سبد خرید خالی است.")
            return
        ORDER_COUNTER += 1
        order_number = ORDER_COUNTER

    total = 0
    customer_message = (
        "✅ سفارش شما با موفقیت ثبت شد.\n\n"
        f"🔢 شماره سفارش: {order_number}\n"
        f"👤 نام: {name}\n"
        f"📞 تماس: {phone}\n\n"
        "📦 اقلام سفارش:\n"
    )
    admin_message = (
        "🚨 سفارش جدید دریافت شد 🚨\n\n"
        f"🔢 شماره سفارش: {order_number}\n"
        f"👤 نام مشتری: {name}\n"
        f"📞 شماره تماس: {phone}\n"
        f"🆔 شناسه تلگرام: {chat_id}\n\n"
        "📦 اقلام سفارش:\n"
    )

    for item in cart:
        item_total = item["price"] * item["quantity"]
        total += item_total
        line = (
            f"\n• {item['name']}\n"
            f"کد: {item['code']}\n"
            f"تعداد: {item['quantity']}\n"
            f"مبلغ: {item_total:,} ریال\n"
        )
        customer_message += line
        admin_message += line

    customer_message += f"\n💰 جمع کل: {total:,} ریال"
    admin_message += f"\n💰 جمع کل: {total:,} ریال"

    send_message(chat_id, customer_message)

    if ADMIN_CHAT_ID:
        send_message(ADMIN_CHAT_ID, admin_message)
    else:
        print("TELEGRAM_ADMIN_ID is not configured; admin notification skipped.")

    with LOCK:
        CARTS[chat_id] = []
        USER_STATES.pop(chat_id, None)

    main_menu(chat_id)


# =========================================================
# CALLBACKS
# =========================================================

def process_callback(callback):
    callback_id = callback.get("id")
    data = callback.get("data", "")
    message = callback.get("message") or {}
    chat = message.get("chat") or {}
    chat_id = chat.get("id")

    if chat_id is None:
        answer_callback(callback_id)
        return

    answer_callback(callback_id)

    if data.startswith("ADD:"):
        try:
            index = int(data.split(":", 1)[1])
        except Exception:
            send_message(chat_id, "❌ دکمه نامعتبر است.")
            return

        state = USER_STATES.get(chat_id, {})
        results = state.get("results", [])
        if index < 0 or index >= len(results):
            send_message(chat_id, "❌ این دکمه منقضی شده است. دوباره جستجو کنید.")
            return

        item = results[index]
        USER_STATES[chat_id] = {
            "state": "quantity",
            "item": item,
            "results": results,
        }
        send_message(
            chat_id,
            f"📦 {item['name']}\n\n"
            "لطفاً تعداد مورد نظر را وارد کنید:\n"
            "مثلاً: 2",
        )


# =========================================================
# MESSAGES
# =========================================================

def process_message(message):
    chat = message.get("chat") or {}
    chat_id = chat.get("id")
    text = (message.get("text") or "").strip()

    if chat_id is None:
        return

    if text.startswith("/start"):
        USER_STATES.pop(chat_id, None)
        send_message(
            chat_id,
            "👋 به ربات تکنو یدک پارت عباسی خوش آمدید.\n\n"
            "برای شروع، یکی از گزینه‌های زیر را انتخاب کنید.",
        )
        main_menu(chat_id)
        return

    if text == "/id":
        send_message(chat_id, f"🆔 شناسه تلگرام شما:\n{chat_id}")
        return

    if text == "📞 تماس مستقیم":
        send_message(
            chat_id,
            f"📞 شماره تماس:\n{PHONE}\n\n"
            "برای تماس مستقیم از همین شماره استفاده کنید.",
        )
        return

    if text == "🌐 ورود به سایت":
        send_message(chat_id, WEBSITE)
        return

    if text == "📄 دریافت لیست قیمت":
        send_message(
            chat_id,
            "📄 برای دریافت لیست قیمت، یکی از گروه‌های زیر را انتخاب کنید.",
        )
        return

    if text == "🔙 منوی اصلی":
        USER_STATES.pop(chat_id, None)
        main_menu(chat_id)
        return

    if text == "🔎 جستجوی کالا":
        USER_STATES[chat_id] = {"state": "search"}
        send_message(
            chat_id,
            "🔎 کد یا نام کالا را ارسال کنید.\n\n"
            "مثلاً: سوکت یا 12345",
        )
        return

    if text == "🛒 سبد خرید":
        USER_STATES.pop(chat_id, None)
        show_cart(chat_id)
        return

    if text == "🗑️ خالی کردن سبد":
        clear_cart(chat_id)
        return

    if text == "🧾 ثبت سفارش":
        start_order(chat_id)
        return

    state_data = USER_STATES.get(chat_id)

    if state_data:
        state = state_data.get("state")

        if state == "search":
            if not text:
                send_message(chat_id, "❌ عبارت جستجو را وارد کنید.")
                return
            send_message(chat_id, "🔎 در حال جستجو...")
            send_search_results(chat_id, search_excel(text))
            return

        if state == "quantity":
            number_text = normalize(text)
            if not number_text.isdigit():
                send_message(
                    chat_id,
                    "❌ تعداد باید عدد باشد.\nمثلاً: 2",
                )
                return

            quantity = int(number_text)
            if quantity <= 0:
                send_message(chat_id, "❌ تعداد باید بیشتر از صفر باشد.")
                return

            item = state_data.get("item")
            results = state_data.get("results", [])

            if not item:
                USER_STATES.pop(chat_id, None)
                send_message(chat_id, "❌ اطلاعات کالا پیدا نشد.")
                return

            add_to_cart(chat_id, item, quantity)
            send_message(
                chat_id,
                f"✅ {item['name']}\n"
                f"به تعداد {quantity} عدد به سبد خرید اضافه شد. 🛒",
            )

            USER_STATES[chat_id] = {
                "state": "search_results",
                "results": results,
            }

            send_message(
                chat_id,
                "👇 می‌توانید کالای دیگری اضافه کنید یا سبد خرید را ببینید.",
                [
                    ["🛒 سبد خرید"],
                    ["🔎 جستجوی کالا"],
                    ["🧾 ثبت سفارش"],
                    ["🔙 منوی اصلی"],
                ],
            )
            return

        if state == "order_name":
            if not text:
                send_message(chat_id, "❌ نام را وارد کنید.")
                return
            state_data["name"] = text
            state_data["state"] = "order_phone"
            send_message(chat_id, "📞 لطفاً شماره تماس خود را وارد کنید:")
            return

        if state == "order_phone":
            phone = re.sub(r"\D", "", normalize(text))
            if len(phone) < 10:
                send_message(
                    chat_id,
                    "❌ شماره تماس صحیح نیست.\nلطفاً دوباره وارد کنید:",
                )
                return
            finish_order(chat_id, state_data["name"], phone)
            return

    if text in PDF_GROUPS:
        USER_STATES.pop(chat_id, None)
        group = PDF_GROUPS[text]
        send_message(chat_id, "⏳ در حال آماده‌سازی فایل PDF...")
        pdf_path = find_pdf(group)

        if not pdf_path:
            send_message(chat_id, "❌ فایل PDF این گروه پیدا نشد.")
            return

        response = send_pdf(chat_id, pdf_path, f"📄 لیست قیمت {group}")
        try:
            if response is not None and response.json().get("ok"):
                return
        except Exception:
            pass

        send_message(chat_id, "❌ ارسال فایل PDF انجام نشد.")
        return

    if text:
        USER_STATES[chat_id] = {"state": "search"}
        send_message(chat_id, "🔎 در حال جستجوی کامل...")
        send_search_results(chat_id, search_excel(text))


# =========================================================
# WEBHOOK / HEALTH
# =========================================================

@app.route("/telegram_webhook", methods=["POST"])
def telegram_webhook():
    try:
        update = request.get_json(silent=True) or {}
        print("TELEGRAM UPDATE:", update)

        if update.get("callback_query"):
            threading.Thread(
                target=process_callback,
                args=(update["callback_query"],),
                daemon=True,
            ).start()
        elif update.get("message"):
            threading.Thread(
                target=process_message,
                args=(update["message"],),
                daemon=True,
            ).start()

        return jsonify({"ok": True})
    except Exception as e:
        print("WEBHOOK ERROR:", repr(e))
        return jsonify({"ok": True})


@app.route("/")
def home():
    return "Telegram Bot is running"


@app.route("/health")
def health():
    return jsonify({
        "status": "ok",
        "telegram_token_configured": bool(TOKEN),
        "admin_configured": bool(ADMIN_CHAT_ID),
    })


def setup_webhook():
    if not TOKEN:
        print("❌ TELEGRAM_TOKEN NOT FOUND")
        return

    render_url = os.getenv("RENDER_EXTERNAL_URL", "").strip()
    if not render_url:
        print("⚠️ RENDER_EXTERNAL_URL NOT FOUND")
        return

    webhook_url = render_url.rstrip("/") + "/telegram_webhook"

    try:
        response = requests.post(
            f"{BASE_URL}/setWebhook",
            json={
                "url": webhook_url,
                "allowed_updates": ["message", "callback_query"],
            },
            timeout=30,
        )
        print("SET TELEGRAM WEBHOOK:", response.status_code, response.text)
    except Exception as e:
        print("SET TELEGRAM WEBHOOK ERROR:", repr(e))


if __name__ == "__main__":
    print("========================================")
    print("TELEGRAM BOT - TECHNO YADAK PART ABBASI")
    print("SEARCH + PDF + CART + ORDER")
    print("========================================")

    def startup():
        time.sleep(3)
        setup_webhook()

    threading.Thread(target=startup, daemon=True).start()

    port = int(os.environ.get("PORT", 10000))
    app.run(host="0.0.0.0", port=port, threaded=True)
